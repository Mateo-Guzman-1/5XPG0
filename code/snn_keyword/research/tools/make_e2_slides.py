"""Five slides on the network-size result (E2) for about one minute of a larger talk.

Builds research/reports/E2_network_size_slides.pptx from research/results/E2/summary.json:
native PowerPoint charts and shapes in theme colors (editable; pasted into another deck
with "Use destination theme", they take that deck's colors and fonts), titles in title
placeholders, and the spoken text in the speaker notes. No introduction or conclusion
slide: the five slides are one section of a larger presentation. Processing time is the
worst case of the plain RV32IM firmware, the same variant on every slide.
Needs python-pptx (pip install python-pptx).

    make_e2_slides.py [--results research/results/E2] [--out research/reports/E2_network_size_slides.pptx]
"""
import argparse
import io
import json
from pathlib import Path
import sys
import uuid
import numpy as np
from lxml import etree
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

TOOLS = Path(__file__).resolve().parent
KWS = TOOLS.parents[1]   # code/snn_keyword
sys.path.insert(0, str(KWS))
from benchmark_inputs import VECTORS

MODEL_LIMIT, BUDGET, CLOCK = 144 * 1024, 25_000_000, 100e6
THEME = {'name': 'E2 network size', 'head': 'Calibri', 'body': 'Calibri',
         'colors': {'dk1': '1D2329', 'lt1': 'FFFFFF', 'dk2': '2E3A46', 'lt2': 'F1F4F6',
                    # Categorical slots of the dataviz reference palette, then the limit red and a gray.
                    'accent1': '2A78D6', 'accent2': 'EB6834', 'accent3': '1BAF7A', 'accent4': 'EDA100',
                    'accent5': 'D03B3B', 'accent6': '898781', 'hlink': '2A78D6', 'folHlink': '52514E'}}
T = MSO_THEME_COLOR
INK, INK2, PANEL = T.TEXT_1, T.TEXT_2, T.BACKGROUND_2
SERIES = [T.ACCENT_1, T.ACCENT_2, T.ACCENT_3]
LIMIT, GRAY = T.ACCENT_5, T.ACCENT_6
GRID, AXIS = 'E1E0D9', 'C3C2B7'   # hex only: chart gridlines and axis lines
SPECTROGRAM = ['FFFFFF', 'CDE2FB', '9EC5F4', '6DA7EC', '3987E5', '256ABF', '184F95', '0D366B']
WIDE_W, WIDE_H = Emu(12192000), Emu(6858000)   # 13.333 x 7.5 in, PowerPoint's 16:9


# ---------------------------------------------------------------- deck and theme

def apply_theme(prs):
    """Theme colors and fonts into theme1.xml; left-aligned titles; master scaled to 16:9."""
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    root.set('name', THEME['name'])
    scheme = root.find(f'.//{{{a}}}clrScheme')
    scheme.set('name', THEME['name'])
    for slot in scheme:
        for child in list(slot):
            slot.remove(child)
        etree.SubElement(slot, f'{{{a}}}srgbClr', val=THEME['colors'][etree.QName(slot).localname])
    fonts = root.find(f'.//{{{a}}}fontScheme')
    fonts.set('name', THEME['name'])
    fonts.find(f'{{{a}}}majorFont/{{{a}}}latin').set('typeface', THEME['head'])
    fonts.find(f'{{{a}}}minorFont/{{{a}}}latin').set('typeface', THEME['body'])
    part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    # Title style: 36 pt bold, left aligned, dark slate (tx2).
    title = prs.slide_master.element.find(f'{qn("p:txStyles")}/{qn("p:titleStyle")}/{qn("a:lvl1pPr")}')
    title.set('algn', 'l')
    rpr = title.find(qn('a:defRPr'))
    rpr.set('sz', '3600')
    rpr.set('b', '1')
    for fill in rpr.findall(qn('a:solidFill')):
        rpr.remove(fill)
    fill = etree.Element(qn('a:solidFill'))
    etree.SubElement(fill, qn('a:schemeClr'), val='tx2')
    rpr.insert(0, fill)
    # The default template is 4:3: stretch every placeholder of master and layouts to 16:9.
    scale = WIDE_W / prs.slide_width
    for holder in [prs.slide_master, *prs.slide_layouts]:
        for ph in holder.placeholders:
            if ph.left is not None and ph.width is not None:
                ph.left, ph.width = Emu(int(ph.left * scale)), Emu(int(ph.width * scale))
    prs.slide_width, prs.slide_height = WIDE_W, WIDE_H
    # Content titles: same box on the master and on every layout that sets its own.
    for holder in [prs.slide_master, *prs.slide_layouts]:
        for ph in holder.placeholders:
            if ph.placeholder_format.type == 1 and (holder is prs.slide_master or ph.left is not None):
                ph.left, ph.top, ph.width, ph.height = Inches(0.6), Inches(0.4), Inches(12.13), Inches(0.9)


def add_section(prs, name):
    """One PowerPoint section around all slides, so the block stays together when merged."""
    p14 = 'http://schemas.microsoft.com/office/powerpoint/2010/main'
    pres = prs.part._element
    ext_lst = pres.find(qn('p:extLst'))
    if ext_lst is None:
        ext_lst = etree.SubElement(pres, qn('p:extLst'))
    ext = etree.SubElement(ext_lst, qn('p:ext'), uri='{521415D9-36F7-43E2-AB2F-B90AF26B5E84}')
    sections = etree.SubElement(ext, f'{{{p14}}}sectionLst', nsmap={'p14': p14})
    section = etree.SubElement(sections, f'{{{p14}}}section', name=name, id='{' + str(uuid.uuid4()).upper() + '}')
    ids = etree.SubElement(section, f'{{{p14}}}sldIdLst')
    for sld in pres.find(qn('p:sldIdLst')):
        etree.SubElement(ids, f'{{{p14}}}sldId', id=sld.get('id'))


def new_slide(prs, title, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[5])   # "Title Only"
    slide.shapes.title.text = title
    slide.notes_slide.notes_text_frame.text = notes
    return slide


# ---------------------------------------------------------------- shapes

def color(target, theme, brightness=0.0):
    target.theme_color = theme
    if brightness:
        target.brightness = brightness


def text(slide, x, y, w, h, runs, size=14, theme=INK, brightness=0.0, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, name=None, spacing=0):
    """Text box; runs = str, or list of paragraphs, each a str or a list of (text, options) runs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    paragraphs = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paragraphs):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        pieces = [(para, {})] if isinstance(para, str) else para
        before = pieces[0][1].get('before', spacing)   # the first run's 'before' sets the paragraph spacing
        if before and i:
            p.space_before = Pt(before)
        for content, opts in pieces:
            run = p.add_run()
            run.text = content
            run.font.size = Pt(opts.get('size', size))
            run.font.bold = opts.get('bold', bold)
            run.font.italic = opts.get('italic', False)
            color(run.font.color, opts.get('theme', theme), opts.get('brightness', brightness))
    return box


def card(slide, x, y, w, h, theme=PANEL, brightness=0.0, name=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = min(0.12, 0.18 / min(w, h))
    shape.fill.solid()
    color(shape.fill.fore_color, theme, brightness)
    shape.line.fill.background()
    shape.shadow.inherit = False
    if name:
        shape.name = name
    return shape


def dot(slide, cx, cy, d, theme, brightness=0.0, name=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    shape.fill.solid()
    color(shape.fill.fore_color, theme, brightness)
    shape.line.fill.background()
    shape.shadow.inherit = False
    if name:
        shape.name = name
    return shape


def line(slide, x1, y1, x2, y2, theme=INK2, brightness=0.75, width=0.75):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    color(connector.line.color, theme, brightness)
    connector.line.width = Pt(width)
    return connector


def stat(slide, x, y, w, h, value, label, key=None, value_size=36, value_w=None, name=None):
    """Rounded card with a big number and a short label; key = series color shown as a dot.
    value_w set: number on the left and label on the right (for wide, low cards)."""
    card(slide, x, y, w, h, name=name)
    pad = 0.3
    vx = x + pad + (0.32 if key is not None else 0)
    if key is not None:
        dot(slide, x + pad + 0.09, y + (h / 2 if value_w else 0.25 + 0.3), 0.18, key)
    if value_w:
        text(slide, vx, y, value_w, h, value, size=value_size, theme=INK2, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, vx + value_w + 0.2, y, x + w - pad - (vx + value_w + 0.2), h, label, size=14, theme=INK,
             brightness=0.25, anchor=MSO_ANCHOR.MIDDLE)
    else:
        text(slide, vx, y + 0.25, x + w - pad - vx, 0.6, value, size=value_size, theme=INK2, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
        text(slide, x + pad, y + 0.93, w - 2 * pad, h - 0.93 - 0.15, label, size=14, theme=INK, brightness=0.25)


def footnote(slide, words, y=6.85):
    text(slide, 0.6, y, 12.13, 0.3, words, size=11, theme=INK, brightness=0.4, name='Footnote')


# ---------------------------------------------------------------- charts

def xy_chart(slide, x, y, w, h, series, ymax, ymajor, yfmt, ytitle, legend=XL_LEGEND_POSITION.TOP):
    """Native scatter chart with straight lines: series = dicts with name, x, y, theme and optional
    sd (error bars), dash, marker, labels {point index: text}."""
    data = XyChartData()
    for s in series:
        ser = data.add_series(s['name'])
        for xv, yv in zip(s['x'], s['y']):
            ser.add_data_point(xv, yv)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, Inches(x), Inches(y), Inches(w), Inches(h), data)
    frame.name = 'Chart'
    chart = frame.chart
    chart.font.size = Pt(12)
    chart.font.name = '+mn-lt'
    color(chart.font.color, INK, 0.3)
    chart.has_legend = legend is not None
    if legend is not None:
        chart.legend.position = legend
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
    xa, ya = chart.category_axis, chart.value_axis   # both are value axes in a scatter chart
    for axis, lo, hi, major, fmt, title in ((xa, 0, 104, 16, '0', 'hidden neurons'),
                                            (ya, 0, ymax, ymajor, yfmt, ytitle)):
        axis.minimum_scale, axis.maximum_scale, axis.major_unit = lo, hi, major
        axis.major_tick_mark = axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.tick_labels.number_format, axis.tick_labels.number_format_is_linked = fmt, False
        axis.tick_labels.font.size = Pt(12)
        axis.has_title = True
        axis.axis_title.text_frame.text = title
        run = axis.axis_title.text_frame.paragraphs[0].runs[0]
        run.font.size, run.font.bold = Pt(12), False
        color(run.font.color, INK, 0.3)
    xa.has_major_gridlines = False
    xa.format.line.color.rgb = RGBColor.from_string(AXIS)
    ya.has_major_gridlines = True
    ya.major_gridlines.format.line.color.rgb = RGBColor.from_string(GRID)
    ya.major_gridlines.format.line.width = Pt(0.75)
    ya.format.line.fill.background()
    for s, ser in zip(series, chart.plots[0].series):
        ser.smooth = False
        color(ser.format.line.color, s['theme'])
        ser.format.line.width = Pt(s.get('width', 2.5))
        if s.get('dash'):
            ser.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if s.get('marker', XL_MARKER_STYLE.CIRCLE) is XL_MARKER_STYLE.NONE:
            ser.marker.style = XL_MARKER_STYLE.NONE
        else:
            ser.marker.style = s.get('marker', XL_MARKER_STYLE.CIRCLE)
            ser.marker.size = 9
            ser.marker.format.fill.solid()
            color(ser.marker.format.fill.fore_color, s['theme'])
            ser.marker.format.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            ser.marker.format.line.width = Pt(1.5)
        for i, label in s.get('labels', {}).items():
            label, position = label if isinstance(label, tuple) else (label, XL_LABEL_POSITION.RIGHT)
            dl = ser.points[i].data_label
            dl.text_frame.text = label
            dl.position = position
            run = dl.text_frame.paragraphs[0].runs[0]
            run.font.size, run.font.bold = Pt(13), True
            color(run.font.color, INK2)
        if 'sd' in s:
            error_bars(ser, s['sd'], s['theme'])
    transparent(chart)
    return chart


def error_bars(series, sd, theme):
    """Custom +-sd error bars (python-pptx has no API for them); before c:xVal, as the schema orders."""
    names = {T.ACCENT_1: 'accent1', T.ACCENT_2: 'accent2', T.ACCENT_3: 'accent3'}
    bars = etree.Element(qn('c:errBars'))
    etree.SubElement(bars, qn('c:errDir'), val='y')
    etree.SubElement(bars, qn('c:errBarType'), val='both')
    etree.SubElement(bars, qn('c:errValType'), val='cust')
    etree.SubElement(bars, qn('c:noEndCap'), val='0')
    for tag in ('c:plus', 'c:minus'):
        literal = etree.SubElement(etree.SubElement(bars, qn(tag)), qn('c:numLit'))
        etree.SubElement(literal, qn('c:formatCode')).text = 'General'
        etree.SubElement(literal, qn('c:ptCount'), val=str(len(sd)))
        for i, v in enumerate(sd):
            etree.SubElement(etree.SubElement(literal, qn('c:pt'), idx=str(i)), qn('c:v')).text = f'{float(v):.6f}'
    ln = etree.SubElement(etree.SubElement(bars, qn('c:spPr')), qn('a:ln'), w=str(Pt(1.5)))
    etree.SubElement(etree.SubElement(ln, qn('a:solidFill')), qn('a:schemeClr'), val=names[theme])
    series._element.find(qn('c:xVal')).addprevious(bars)


def transparent(chart):
    """No fill or border on chart and plot area: the chart sits directly on the slide."""
    space = chart._chartSpace
    plot = space.find(qn('c:chart')).find(qn('c:plotArea'))
    for parent, anchor in ((space, space.find(qn('c:chart'))), (plot, None)):
        sp = parent.find(qn('c:spPr'))
        if sp is None:
            sp = etree.Element(qn('c:spPr'))
            if anchor is not None:
                anchor.addnext(sp)
            else:
                ext = parent.find(qn('c:extLst'))
                ext.addprevious(sp) if ext is not None else parent.append(sp)
        for child in list(sp):
            sp.remove(child)
        etree.SubElement(sp, qn('a:noFill'))
        etree.SubElement(etree.SubElement(sp, qn('a:ln')), qn('a:noFill'))
    rounded = space.find(qn('c:roundedCorners'))
    if rounded is not None:
        rounded.set('val', '0')


# ---------------------------------------------------------------- slides

def spectrogram_png():
    """The real 24 x 32 input of a "Sheila" test clip (verification vector 0), as a small PNG."""
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.colors import LinearSegmentedColormap
    x = np.load(VECTORS)['x'][0].reshape(24, 32)
    fig = plt.figure(figsize=(3.2, 2.4), dpi=200)
    ax = fig.add_axes((0, 0, 1, 1))
    ax.imshow(x, origin='lower', aspect='auto', interpolation='nearest',
              cmap=LinearSegmentedColormap.from_list('blues', ['#' + c for c in SPECTROGRAM]))
    ax.set_axis_off()
    buf = io.BytesIO()
    fig.savefig(buf, format='png')
    plt.close(fig)
    buf.seek(0)
    return buf


def slide_network(prs, sizes):
    s = new_slide(prs, 'Network size = number of hidden neurons',
                  'Here we changed one thing only: the number of hidden neurons, from 8 up to 95, the most '
                  'that fits. Each size was trained three times; everything else stayed the same.')
    # Diagram: input picture -> hidden neurons -> two outputs.
    picture = s.shapes.add_picture(spectrogram_png(), Inches(0.6), Inches(2.45), Inches(2.6), Inches(1.95))
    picture.name = 'Input picture'
    color(picture.line.color, INK2, 0.6)   # thin frame: the picture's white background has no edge
    picture.line.width = Pt(0.75)
    text(s, 0.6, 4.55, 2.6, 0.4, '768 inputs', size=16, bold=True, theme=INK2)
    text(s, 0.6, 4.95, 2.6, 0.6, ['1 s sound picture,', '24 pitches × 32 time steps'], size=12, theme=INK,
         brightness=0.3)
    hx, hd, ox, od = 4.55, 0.42, 6.35, 0.5
    hidden = [1.75, 2.35, 2.95, 3.55, 4.15, 4.75, 5.75]
    outputs = {'Sheila': 3.05, 'not Sheila': 4.45}
    for hy in hidden:
        for iy in (2.65, 3.15, 3.65, 4.15):
            line(s, 3.25, iy, hx - hd / 2, hy)
        for oy in outputs.values():
            line(s, hx + hd / 2, hy, ox - od / 2, oy)
    for hy in hidden:
        dot(s, hx, hy, hd, SERIES[0], name='Hidden neuron')
    for k in range(3):
        dot(s, hx, 5.12 + 0.13 * k, 0.07, INK2, 0.3)
    text(s, hx - 1.2, 6.15, 2.4, 0.4, 'H hidden neurons', size=16, bold=True, theme=INK2, align=PP_ALIGN.CENTER)
    text(s, hx - 1.2, 6.53, 2.4, 0.35, '= the network size', size=12, theme=INK, brightness=0.3, align=PP_ALIGN.CENTER)
    for name, oy in outputs.items():
        dot(s, ox, oy, od, SERIES[1], name='Output')
        text(s, ox + 0.38, oy - 0.2, 1.4, 0.4, name, size=14, theme=INK2, anchor=MSO_ANCHOR.MIDDLE)
    # What was tested.
    x0, y0, w = 8.45, 1.6, 4.28
    card(s, x0, y0, w, 5.1, name='Sizes tested')
    text(s, x0 + 0.35, y0 + 0.3, w - 0.7, 0.4, 'Sizes tested', size=16, bold=True, theme=INK2)
    chip_w, gap = 0.8, 0.13
    for i, h in enumerate(sizes):
        row, col = divmod(i, 4)
        cx, cy = x0 + 0.35 + col * (chip_w + gap), y0 + 0.85 + row * 0.62
        last = h == sizes[-1]
        chip = card(s, cx, cy, chip_w, 0.48, theme=SERIES[0] if last else T.BACKGROUND_1, name=f'Size {h}')
        chip.adjustments[0] = 0.5
        frame = chip.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(h)
        run.font.size, run.font.bold = Pt(16), True
        color(run.font.color, T.BACKGROUND_1 if last else INK2)
    text(s, x0 + 0.35, y0 + 2.05, w - 0.7, 0.35, '95 = the most that fits in memory', size=12, theme=INK,
         brightness=0.3)
    text(s, x0 + 0.35, y0 + 2.65, w - 0.7, 2.25,
         [[('3 training runs per size', {'bold': True, 'theme': INK2})],
          [('results are the mean ± the spread', {'size': 13, 'brightness': 0.3, 'before': 2})],
          [('Everything else identical', {'bold': True, 'theme': INK2, 'before': 14})],
          [('same data, training and threshold rules', {'size': 13, 'brightness': 0.3, 'before': 2})],
          [(f'All {3 * len(sizes)} models run on the board', {'bold': True, 'theme': INK2, 'before': 14})],
          [('results identical to the Python reference', {'size': 13, 'brightness': 0.3, 'before': 2})]],
         size=16)


def slide_cost(prs, rows):
    h = [r['hidden'] for r in rows]
    memory = [r['model_bytes'] / MODEL_LIMIT for r in rows]
    time = [r['base_cycles_max'] / BUDGET for r in rows]
    per_neuron_ms = np.polyfit([r['hidden'] for r in rows if r['hidden'] >= 32],
                               [r['base_cycles_max'] / CLOCK * 1e3 for r in rows if r['hidden'] >= 32], 1)[0]
    s = new_slide(prs, 'Memory runs out at 95 neurons, time does not',
                  'Each neuron costs one and a half kilobytes and almost a millisecond. At 95 neurons the '
                  'memory is full, but processing uses only a third of the time per window: memory is the limit.')
    # No label on the last memory point: it sits on the limit line, which already says "full".
    xy_chart(s, 0.6, 1.45, 8.1, 5.25, [
        {'name': 'Memory (of 144 KiB)', 'x': h, 'y': memory, 'theme': SERIES[0]},
        {'name': 'Processing time (of 250 ms)', 'x': h, 'y': time, 'theme': SERIES[1],
         'labels': {len(h) - 1: f'{100 * time[-1]:.1f}%'}},
        {'name': 'Limit', 'x': [0, 104], 'y': [1, 1], 'theme': LIMIT, 'dash': True, 'width': 1.75,
         'marker': XL_MARKER_STYLE.NONE}],
        ymax=1.1, ymajor=0.2, yfmt='0%', ytitle='share of the platform limit')
    x0, w = 9.05, 3.68
    largest = rows[-1]
    stat(s, x0, 1.6, w, 1.5, f'{largest["model_bytes"] / largest["hidden"] / 1024:.1f} KiB',
         'memory per neuron', key=SERIES[0], name='Memory per neuron')
    stat(s, x0, 3.4, w, 1.5, f'{per_neuron_ms:.2f} ms', 'processing time per neuron', key=SERIES[1],
         name='Time per neuron')
    stat(s, x0, 5.2, w, 1.5, f'{largest["hidden"]} neurons', f'memory full, time {100 * time[-1]:.1f}% used',
         key=LIMIT, value_size=32, name='Largest size')
    footnote(s, 'Processing time: worst case on the board\'s RISC-V core (plain RV32IM firmware). '
                'Memory: the weights in the 144 KiB model area.')


def slide_quality(prs, rows):
    h = [r['hidden'] for r in rows]
    f1 = [r['test_f1'] for r in rows]
    live = [r['live_2of3_detected'] for r in rows]
    s = new_slide(prs, 'Quality stops improving after 16 neurons',
                  'Quality jumps from 8 to 16 neurons and then stays flat. Above 16, the differences between '
                  'sizes are smaller than the spread between training runs.')
    pct = lambda v: f'{100 * v:.0f}%'
    xy_chart(s, 0.6, 1.45, 12.13, 3.95, [
        # Labels only at both ends, outside the line and clear of the spread bars.
        {'name': 'F1 score, test clips', 'x': h, 'y': f1, 'theme': SERIES[0],
         'sd': [r['test_f1_sd'] for r in rows],
         'labels': {0: (pct(f1[0]), XL_LABEL_POSITION.LEFT), len(h) - 1: pct(f1[-1])}},
        {'name': '"Sheila" detected, live stream (2 of 3)', 'x': h, 'y': live, 'theme': SERIES[1],
         'marker': XL_MARKER_STYLE.SQUARE, 'sd': [r['live_2of3_detected_sd'] for r in rows],
         'labels': {0: (pct(live[0]), XL_LABEL_POSITION.LEFT), len(h) - 1: pct(live[-1])}}],
        ymax=1.0, ymajor=0.2, yfmt='0%', ytitle='score')
    stat(s, 0.6, 5.55, 5.9, 1.15, f'{pct(f1[0])} → {pct(f1[1])}', 'F1 score from 8 to 16 neurons', value_size=30,
         value_w=2.7, name='Steep part')
    stat(s, 6.83, 5.55, 5.9, 1.15, 'flat after 16', 'differences smaller than the spread of 3 training runs',
         value_size=30, value_w=2.7, name='Plateau')
    footnote(s, 'Mean of 3 training runs per size; bars show ± 1 standard deviation. Live stream: 1 s windows '
                'every 250 ms; a detection needs 2 of 3.')


def slide_near_misses(prs, rows):
    h = [r['hidden'] for r in rows]
    s = new_slide(prs, 'More neurons do not fix near misses',
                  'A bigger network also does not reject near misses: about a third of words like Sheena or '
                  'shield still pass, at every size. That needs a network that follows the order of sounds.')
    x0, w = 0.6, 3.9
    card(s, x0, 1.6, w, 5.15, name='Near-miss callout')
    text(s, x0 + 0.35, 1.9, w - 0.7, 0.9, '≈ 1/3', size=48, bold=True, theme=INK2)
    text(s, x0 + 0.35, 2.85, w - 0.7, 0.7, ['of near misses still pass,', 'at every size'], size=16, theme=INK2,
         bold=True)
    text(s, x0 + 0.35, 3.85, w - 0.7, 0.75,
         'Near misses sound like "Sheila": Sheena, shield, "shee" + another word.', size=14, theme=INK, brightness=0.25)
    text(s, x0 + 0.35, 4.85, w - 0.7, 0.75,
         'One dense layer sees the whole second at once, not the order of the sounds.',
         size=14, theme=INK, brightness=0.25)
    text(s, x0 + 0.35, 5.85, w - 0.7, 0.65, '→ needs a network that follows the order of sounds', size=14,
         theme=INK2, bold=True)
    near = [('"shee" + another word', 'hardneg_ending', SERIES[0], XL_MARKER_STYLE.CIRCLE),
            ('she-, sheet, shield, Sheena (computer voices)', 'tts_2of3_she- words', SERIES[1], XL_MARKER_STYLE.SQUARE),
            ('another start + "eila"', 'hardneg_onset', SERIES[2], XL_MARKER_STYLE.TRIANGLE),
            ('"sh" alone', 'hardneg_sh', GRAY, XL_MARKER_STYLE.DIAMOND)]
    xy_chart(s, 4.85, 1.45, 7.88, 5.25, [
        {'name': name, 'x': h, 'y': [r[key] for r in rows], 'theme': theme, 'marker': marker}
        for name, key, theme, marker in near],
        ymax=0.6, ymajor=0.1, yfmt='0%', ytitle='share accepted as "Sheila"', legend=XL_LEGEND_POSITION.BOTTOM)
    footnote(s, 'Accepted at the live threshold: per 1 s window (edited real recordings) or per word with "2 of 3" '
                '(computer voices). Mean of 3 runs.')


def slide_efficient(prs, rows):
    by = {r['hidden']: r for r in rows}
    s = new_slide(prs, '48 neurons match 95 with half the memory',
                  'So 48 neurons give the same quality as 95 with half the memory and half the processing '
                  'time: the efficient size for this design.')
    columns = [(48, 'efficient size', 0.6, SERIES[0], 0.85), (95, 'deployed now', 6.88, PANEL, 0.0)]
    for h, role, x0, theme, brightness in columns:
        r, w = by[h], 5.85
        card(s, x0, 1.6, w, 4.95, theme=theme, brightness=brightness, name=f'{h} neurons')
        text(s, x0 + 0.4, 1.85, w - 0.8, 0.6, f'{h} neurons', size=32, bold=True, theme=INK2)
        text(s, x0 + 0.4, 2.45, w - 0.8, 0.4, role, size=16, theme=INK, brightness=0.25)
        values = [('Memory', f'{r["model_bytes"] / 1024:.0f} KiB ({100 * r["model_bytes"] / MODEL_LIMIT:.1f}%)'),
                  ('Processing time', f'{r["base_cycles_max"] / CLOCK * 1e3:.0f} ms'),
                  ('F1 score', f'{100 * r["test_f1"]:.0f}%'),
                  ('Live detection', f'{100 * r["live_2of3_detected"]:.0f}%')]
        for i, (label, value) in enumerate(values):
            yy = 3.15 + i * 0.8
            text(s, x0 + 0.4, yy, 2.4, 0.6, label, size=16, theme=INK, brightness=0.25, anchor=MSO_ANCHOR.MIDDLE)
            text(s, x0 + 2.6, yy, w - 3.0, 0.6, value, size=28, bold=True, theme=INK2, align=PP_ALIGN.RIGHT,
                 anchor=MSO_ANCHOR.MIDDLE)
    footnote(s, '48 has the best mean validation F1 (chosen before looking at test results); its quality '
                'differences to 95 are within the spread of 3 runs.')


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument('--results', type=Path, default=TOOLS.parent / 'results/E2')
    p.add_argument('--out', type=Path, default=TOOLS.parent / 'reports/E2_network_size_slides.pptx')
    a = p.parse_args()
    rows = json.loads((a.results / 'summary.json').read_text())['summary']
    prs = Presentation()
    apply_theme(prs)
    props = prs.core_properties
    props.title, props.subject = 'Network size: memory, time and quality', 'Research experiment E2 (5XPG0, Group 2)'
    props.author = props.last_modified_by = '5XPG0 Group 2'
    slide_network(prs, [r['hidden'] for r in rows])
    slide_cost(prs, rows)
    slide_quality(prs, rows)
    slide_near_misses(prs, rows)
    slide_efficient(prs, rows)
    add_section(prs, 'Network size (E2)')
    a.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(a.out)
    print(f'Wrote {a.out} ({len(prs.slides)} slides)')


if __name__ == '__main__':
    main()
